# -*- coding: utf-8 -*-
import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fp = r"C:\Users\nikan\Documents\Společenství\Výpisy\Vypisy_2026_Komplet.xlsx"
wb = openpyxl.load_workbook(fp, data_only=True)

print("Searching for any transaction with ~24 359 or account 464441023 across ALL sheets...")
for sname in wb.sheetnames:
    if not sname.startswith('Výpis'): continue
    ws = wb[sname]
    for r in range(12, ws.max_row + 1):
        vydaj = ws.cell(r, 12).value
        prijem = ws.cell(r, 11).value
        protiucet = str(ws.cell(r, 5).value or '')
        popis = str(ws.cell(r, 7).value or '')
        datum = ws.cell(r, 2).value
        
        # Check against target amounts or account
        if '464441023' in protiucet or 'večeř' in popis.lower():
            print(f"MATCH ACC/NAME: {sname} | {datum} | Příjem: {prijem} | Výdaj: {vydaj} | Protiúčet: {protiucet} | Popis: {popis}")
        
        if vydaj is not None and (24000 <= abs(vydaj) <= 25500):
            print(f"MATCH AMOUNT: {sname} | {datum} | Výdaj: {vydaj} | Protiúčet: {protiucet} | Popis: {popis}")
