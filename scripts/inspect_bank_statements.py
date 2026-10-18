# -*- coding: utf-8 -*-
import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fp = r"C:\Users\nikan\Documents\Společenství\Výpisy\Vypisy_2026_Komplet.xlsx"
wb = openpyxl.load_workbook(fp, data_only=True)
print("Sheets:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n--- Sheet: {sname} (max row: {ws.max_row}, max col: {ws.max_column}) ---")
    for r in range(1, min(15, ws.max_row + 1)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
        if any(v is not None for v in row_vals):
            print(f"R{r:02d}: {row_vals}")
