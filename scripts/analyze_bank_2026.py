# -*- coding: utf-8 -*-
import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fp = r"C:\Users\nikan\Documents\Společenství\Výpisy\Vypisy_2026_Komplet.xlsx"
wb = openpyxl.load_workbook(fp, data_only=True)

all_tx = []

for sname in wb.sheetnames:
    if not sname.startswith("Výpis"):
        continue
    ws = wb[sname]
    # Header is at row 11
    # Data starts at row 12
    for r in range(12, ws.max_row + 1):
        p_c = ws.cell(r, 1).value
        datum = ws.cell(r, 2).value
        typ = ws.cell(r, 4).value
        protiucet = ws.cell(r, 5).value
        nazev = ws.cell(r, 6).value
        popis = ws.cell(r, 7).value
        vs = ws.cell(r, 8).value
        prijem = ws.cell(r, 11).value
        vydaj = ws.cell(r, 12).value
        castka = ws.cell(r, 13).value
        zustatek = ws.cell(r, 14).value
        
        if datum is None and castka is None:
            continue
        
        all_tx.append({
            'sheet': sname,
            'p_c': p_c,
            'datum': str(datum).strip() if datum else '',
            'typ': str(typ).strip() if typ else '',
            'protiucet': str(protiucet).strip() if protiucet else '',
            'nazev': str(nazev).strip() if nazev else '',
            'popis': str(popis).strip() if popis else '',
            'vs': str(vs).strip() if vs else '',
            'prijem': prijem,
            'vydaj': vydaj,
            'castka': castka,
            'zustatek': zustatek
        })

print(f"Total transactions loaded from 2026: {len(all_tx)}")

# 1. Analyze share payouts (výdaje s podíly / výnosy)
print("\n" + "="*80)
print("1. VÝPLATY PODÍLŮ Z VÝNOSŮ (SPOLUVLASTNĚNÉ JEDNOTKY 2026)")
print("="*80)
share_payouts = []
for tx in all_tx:
    # check if expense and matches share payout keywords or names
    txt = (tx['popis'] + " " + tx['nazev']).lower()
    if tx['vydaj'] is not None and (
        any(k in txt for k in ['výnos', 'podíl', 'podil', 'vynos']) or
        any(n in txt for n in ['sebera', 'soldát', 'soldat', 'bertuzzi', 'procházk', 'prochazk', 'andělov', 'andelov', 'pořízk', 'porizk', 'jalový', 'jalovy', 'hábov', 'habov', 'nývlt', 'nyvlt', 'pavlíčk', 'pavlicek', 'rumler'])
    ):
        share_payouts.append(tx)
        print(f"Datum: {tx['datum']} | Částka: {tx['vydaj']:>10} Kč | Příjemce: {tx['nazev'][:30]} | Popis: {tx['popis']}")

# 2. Analyze rental income (příjmy od nájemců)
print("\n" + "="*80)
print("2. PŘÍJMY Z NÁJEMNÉHO (DDAANN, KOPECKÝ, ONARI, BOSCHOFSKY)")
print("="*80)
rents = {'DDAANN': [], 'KOPECKÝ': [], 'ONARI': [], 'BOSCHOFSKY': []}
for tx in all_tx:
    txt = (tx['popis'] + " " + tx['nazev']).lower()
    if tx['prijem'] is not None:
        if 'ddaann' in txt:
            rents['DDAANN'].append(tx)
        elif 'kopeck' in txt:
            rents['KOPECKÝ'].append(tx)
        elif 'onari' in txt:
            rents['ONARI'].append(tx)
        elif 'bosch' in txt:
            rents['BOSCHOFSKY'].append(tx)

for tenant, items in rents.items():
    print(f"\n--- Nájemce: {tenant} (Počet plateb: {len(items)}) ---")
    tot = sum(i['prijem'] for i in items)
    for i in items:
        print(f"  {i['datum']} | {i['prijem']:>8} Kč | VS: {i['vs']} | {i['popis'][:50]}")
    print(f"  CELKEM ZAPLACENO: {tot:>8} Kč")

# 3. Check David Večeřa payments
print("\n" + "="*80)
print("3. PLATBY PANA DAVIDA VEČEŘI (2026)")
print("="*80)
vecera_tx = [tx for tx in all_tx if 'večeř' in (tx['popis'] + " " + tx['nazev']).lower() or tx['vs'] in ['104', '367104']]
if vecera_tx:
    for v in vecera_tx:
        print(f"  {v['datum']} | Příjem: {v['prijem']} | Výdaj: {v['vydaj']} | {v['popis']}")
else:
    print("  ŽÁDNÁ PLATBA OD PANA VEČEŘI NENALEZENA V CELÉM ROCE 2026!")

