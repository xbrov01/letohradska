# -*- coding: utf-8 -*-
import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

fp = r"C:\Users\nikan\Documents\Společenství\Výpisy\Vypisy_2026_Komplet.xlsx"
wb = openpyxl.load_workbook(fp, data_only=True)

all_tx = []
for sname in wb.sheetnames:
    if not sname.startswith("Výpis"):
        continue
    ws = wb[sname]
    for r in range(12, ws.max_row + 1):
        vydaj = ws.cell(r, 12).value
        prijem = ws.cell(r, 11).value
        if vydaj is not None or prijem is not None:
            all_tx.append({
                'datum': ws.cell(r, 2).value,
                'vydaj': vydaj,
                'prijem': prijem,
                'protiucet': ws.cell(r, 5).value,
                'nazev': ws.cell(r, 6).value,
                'popis': ws.cell(r, 7).value,
                'vs': str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value else '',
            })

# List of co-owners from excel
coowners_units = {
    '5': 'Ouředníček J.',
    '6': 'Holubovi',
    '7': 'Horák R.',
    '8': 'Měsíčková I.',
    '10': 'Mandelíková (dědic)',
    '11': 'Zímová J.',
    '12': 'Kopecká I.',
    '13': 'Šoleovi',
    '14': 'Pořízkovi',
    '15': 'Bukačovi',
    '16': 'Soldát R.',
    '17': 'Kitzlerová E.',
    '18': 'Procházkovi',
    '19': 'Rumlerovi',
    '20': 'Bertuzzi G.',
    '21': 'Trefní M.',
    '22': 'Soppé M.',
    '23': 'Sebera I.',
    '24': 'Hábová K. / Andělová B.',
    '25': 'Brožovi',
    '103': 'Kaifosz J.',
    '104': 'Večeřa D.',
}

print("=== CHECKING SHARE PAYOUTS FOR EACH UNIT / CO-OWNER IN 2026 ===")
for vs_key, owner in coowners_units.items():
    # look for payout in tx where vs == vs_key or in popis/nazev and amount > 3000
    payouts = []
    for tx in all_tx:
        if tx['vydaj'] and tx['vydaj'] < -3000:
            txt = (str(tx['popis']) + " " + str(tx['nazev']) + " " + str(tx['protiucet'])).lower()
            if tx['vs'] in [vs_key, '367' + vs_key, '3670' + vs_key] or any(part in txt for part in owner.lower().split() if len(part) > 3):
                # exclude regular supplier invoices like OTIS, PRE, etc.
                if not any(sup in txt for sup in ['otis', 'pre', 'pražsk', 'komi', 'bozp', 'plyn', 'svoz']):
                    payouts.append(tx)
    
    if payouts:
        print(f"[{vs_key:>3}] {owner:<25}: VYPLACENO ({len(payouts)} platba/plateb)")
        for p in payouts:
            print(f"      -> {p['datum']} | {p['vydaj']:>8} Kč | VS: {p['vs']} | Popis: {p['popis']}")
    else:
        print(f"[{vs_key:>3}] {owner:<25}: *** NENALEZENA ŽÁDNÁ VÝPLATA V ROCE 2026 ***")
